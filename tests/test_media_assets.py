"""Synthetic DOCX stress tests for image-only, floating, VML and table images."""
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
import json
import sys
import unittest

from docx import Document
from docx.shared import Inches
from lxml import etree as ET
from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from media_assets import (crop_for_frame, scan_docx, write_inventory, verify_inventory,
                          write_project_matches, verify_project_matches,
                          suggest_photo_layout, suggest_project_photo_pages,
                          validate_asset_bindings, draft_project_catalog)

R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
V = 'urn:schemas-microsoft-com:vml'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def fixture(directory):
    root = Path(directory)
    a = root / 'a.png'; b = root / 'b.png'; c = root / 'c.png'
    first = Image.new('RGB', (240, 160), '#aa3300')
    ImageDraw.Draw(first).ellipse((90, 50, 150, 110), fill='white')
    first.save(a)
    tall = Image.new('RGB', (160, 240), '#003399')
    pen = ImageDraw.Draw(tall)
    pen.rectangle((0, 0, 159, 38), fill='#ffcc00')
    pen.rectangle((0, 202, 159, 239), fill='#ff4444')
    pen.ellipse((55, 95, 105, 145), fill='white')
    tall.save(b)
    wide = Image.new('RGB', (500, 160), '#119955')
    pen = ImageDraw.Draw(wide)
    pen.rectangle((0, 0, 75, 159), fill='#cc22aa')
    pen.rectangle((425, 0, 499, 159), fill='#ffcc00')
    pen.ellipse((220, 50, 280, 110), fill='white')
    wide.save(c)
    doc = Document()
    doc.add_heading('星河项目', level=1)
    doc.add_paragraph('图1：星河项目实施现场')
    image_para = doc.add_paragraph()
    first = image_para.add_run().add_picture(str(a), width=Inches(2))
    a_rid = first._inline.xpath('.//a:blip')[0].get('{' + R + '}embed')
    doc.add_paragraph().add_run().add_picture(str(a), width=Inches(2))
    doc.add_heading('云梯项目', level=1)
    row = doc.add_table(rows=1, cols=1).cell(0, 0)
    row.paragraphs[0].text = '云梯项目课堂照片'
    row.add_paragraph().add_run().add_picture(str(b), width=Inches(2))
    floating = doc.add_paragraph().add_run().add_picture(str(c), width=Inches(2))
    floating._inline.tag = '{' + WP + '}anchor'
    # Reuse one media part via a legacy VML occurrence under a conflicting chapter.
    p = doc.add_paragraph()
    pict = ET.SubElement(p.add_run()._r, '{' + W + '}pict')
    shape = ET.SubElement(pict, '{' + V + '}shape')
    ET.SubElement(shape, '{' + V + '}imagedata').set('{' + R + '}id', a_rid)
    path = root / 'rich.docx'; doc.save(path)
    return path


class MediaAssetsTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.doc = fixture(self.root)
        self.inventory_path = self.root / 'photos.json'
        self.ledger = write_inventory(self.doc, 'speech', self.inventory_path)
        self.projects = {'projects': [
            {'id': 'star', 'name': '星河项目', 'aliases': ['星河']},
            {'id': 'cloud', 'name': '云梯项目', 'aliases': ['云梯']}],
            'reviewed': True,
            'source_doc_sha256': self.ledger['document_sha256']}
        self.catalog_path = self.root / 'projects.json'
        self.catalog_path.write_text(json.dumps(self.projects, ensure_ascii=False), encoding='utf-8')

    def tearDown(self):
        self.temp.cleanup()

    def test_inventory_all_drawing_kinds_and_textless_anchor(self):
        assets = self.ledger['assets']
        self.assertEqual(len(assets), 5)
        self.assertEqual({a['inline_or_anchor'] for a in assets}, {'inline', 'anchor', 'pict'})
        self.assertTrue(all(a['paragraph_anchor'] for a in assets))
        self.assertTrue(all(a['status'] == 'ready' for a in assets))
        self.assertEqual(assets[0]['paragraph_text'], '')
        self.assertEqual(assets[0]['caption_candidates'][0]['text'], '图1：星河项目实施现场')
        self.assertEqual(assets[2]['table']['cell_text'], '云梯项目课堂照片')
        self.assertEqual(assets[0]['media_sha256'], assets[1]['media_sha256'])
        self.assertEqual(assets[0]['media_sha256'], assets[4]['media_sha256'])
        self.assertEqual(len({a['occurrence_id'] for a in assets}), 5)
        self.assertEqual(verify_inventory(self.inventory_path)['assets'], assets)

    def test_project_matching_conflict_and_review_sheet(self):
        draft = draft_project_catalog(self.ledger)
        self.assertEqual([p['name'] for p in draft['projects']], ['星河项目', '云梯项目'])
        self.assertFalse(draft['reviewed'])
        unreviewed = self.root / 'unreviewed.json'
        unreviewed.write_text(json.dumps(draft, ensure_ascii=False), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'reviewed:true'):
            write_project_matches(self.inventory_path, unreviewed, self.root / 'bad-matches.json')
        report_path = self.root / 'matches.json'
        report = write_project_matches(self.inventory_path, self.catalog_path, report_path)
        self.assertEqual(len(report['decisions']), 5)
        self.assertEqual([d['status'] for d in report['decisions']], ['review', 'review', 'auto', 'auto', 'review'])
        self.assertEqual(report['decisions'][2]['project_id'], 'cloud')
        self.assertEqual(report['decisions'][0]['reason'], 'same_media_conflicting_project_context')
        self.assertIn('待确认项目图片', report_path.with_suffix('.review.html').read_text(encoding='utf-8'))
        grouped = suggest_project_photo_pages(self.ledger, report)
        self.assertEqual(len(grouped['projects']), 1)
        self.assertEqual(grouped['projects'][0]['project_id'], 'cloud')
        self.assertEqual(grouped['projects'][0]['pages'][0]['template_slide'], 11)
        self.assertEqual(len(grouped['projects'][0]['asset_ids']), 2)
        self.assertEqual(len(grouped['review_assets']), 3)
        verify_project_matches(report_path)

    def test_layout_crops_and_plan_bindings(self):
        assets = self.ledger['assets']
        self.assertEqual(crop_for_frame([200, 100], [0, 0, 100, 100]),
                         {'l': 25000, 't': 0, 'r': 25000, 'b': 0})
        layout = suggest_photo_layout([assets[2]['asset_id']], self.ledger)
        self.assertEqual(layout['template_slide'], 11)
        self.assertEqual(layout['pictures']['37']['action'], 'remove')
        self.assertEqual(layout['geometry']['34']['bounds_pt'], [30, 150, 876, 240])
        triple = suggest_photo_layout([assets[i]['asset_id'] for i in (0, 2, 3)], self.ledger)
        self.assertEqual(triple['template_slide'], 13)
        self.assertEqual(list(triple['pictures']), ['38', '41', '44'])
        match_path = self.root / 'matches.json'
        write_project_matches(self.inventory_path, self.catalog_path, match_path)
        asset = assets[2]
        plan = {'media_inventory': self.inventory_path.name, 'project_matches': match_path.name,
                'asset_dispositions': {a['asset_id']: {'status': 'used'} if a is asset else
                                       {'status': 'omit', 'reason': '测试未上屏'} for a in assets},
                'pages': [{'id': 'page-one', 'pictures': {'34': {'action': 'replace',
                            'path': asset['extracted_path'], 'sha256': asset['media_sha256'],
                            'asset_id': asset['asset_id'], 'project_id': 'cloud'}}}]}
        result = validate_asset_bindings(plan, self.root)
        self.assertEqual(result['bindings'][0]['occurrence_id'], asset['occurrence_id'])
        self.assertEqual(result['bindings'][0]['attribution'], 'auto')
        self.assertEqual(result['gaps'], [])
        unshown = assets[0]['asset_id']
        plan['asset_dispositions'][unshown] = {'status': 'pending', 'reason': '无法确认图注'}
        self.assertEqual(validate_asset_bindings(plan, self.root)['gaps'][0]['asset_id'], unshown)
        del plan['asset_dispositions'][unshown]
        with self.assertRaisesRegex(ValueError, 'Every DOCX image occurrence'):
            validate_asset_bindings(plan, self.root)
        plan['asset_dispositions'][unshown] = {'status': 'omit', 'reason': '测试未上屏'}
        plan['pages'][0]['pictures']['34']['project_id'] = 'star'
        with self.assertRaisesRegex(ValueError, 'project differs'):
            validate_asset_bindings(plan, self.root)

    def test_reject_changed_document_and_unconfirmed_ambiguous_photo(self):
        match_path = self.root / 'matches.json'
        write_project_matches(self.inventory_path, self.catalog_path, match_path)
        asset = self.ledger['assets'][0]
        plan = {'media_inventory': self.inventory_path.name, 'project_matches': match_path.name,
                'asset_dispositions': {a['asset_id']: {'status': 'used'} if a is asset else
                                       {'status': 'omit', 'reason': '测试未上屏'} for a in self.ledger['assets']},
                'pages': [{'id': 'one', 'pictures': {'34': {'action': 'replace',
                          'path': asset['extracted_path'], 'sha256': asset['media_sha256'],
                          'asset_id': asset['asset_id'], 'project_id': 'star'}}}]}
        with self.assertRaisesRegex(ValueError, 'manual project confirmation'):
            validate_asset_bindings(plan, self.root)
        plan['pages'][0]['pictures']['34']['confirmation_note'] = '人工核对原稿图注与照片内容'
        self.assertEqual(validate_asset_bindings(plan, self.root)['bindings'][0]['attribution'], 'manual')
        self.doc.write_bytes(self.doc.read_bytes() + b'changed')
        with self.assertRaisesRegex(ValueError, 'fingerprint changed'):
            verify_inventory(self.inventory_path)

    def test_external_image_is_reported_and_cannot_enter_photo_layout(self):
        external = self.root / 'external.docx'
        rid = self.ledger['assets'][0]['relationship_id']
        with ZipFile(self.doc) as original, ZipFile(external, 'w', ZIP_DEFLATED) as changed:
            for name in original.namelist():
                data = original.read(name)
                if name == 'word/_rels/document.xml.rels':
                    root = ET.fromstring(data)
                    rel = next(x for x in root if x.get('Id') == rid)
                    rel.set('Target', 'https://example.invalid/photo.png')
                    rel.set('TargetMode', 'External')
                    data = ET.tostring(root)
                changed.writestr(name, data)
        ledger = scan_docx(external, 'speech')
        self.assertEqual(ledger['assets'][0]['status'], 'external_link')
        self.assertEqual(ledger['assets'][0]['external_target'], 'https://example.invalid/photo.png')
        with self.assertRaisesRegex(ValueError, 'extractable JPEG/PNG'):
            suggest_photo_layout([ledger['assets'][0]['asset_id']], ledger)


if __name__ == '__main__':
    unittest.main()
