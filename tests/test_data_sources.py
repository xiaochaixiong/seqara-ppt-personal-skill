from pathlib import Path
from tempfile import TemporaryDirectory
import json
import subprocess
import sys
import unittest
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from data_sources import extract_data, load_data_sources, evaluate_metrics
from native_pptx import TEMPLATE, file_sha, inventory, write_json
from pipeline import extract_source, evaluate_plan_facts, validate_plan


class DataSourceTests(unittest.TestCase):
    def setUp(self):
        temp = TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.base = Path(temp.name)

    def ledger(self, filename, content, sid='table'):
        source = self.base / filename
        source.write_text(content, encoding='utf-8')
        ledger = extract_data(source, sid)
        ledger['path'] = filename
        write_json(self.base / 'data.json', ledger)
        return source, ledger

    def metric(self, mid, value, ref, **other):
        item = {'id': mid, 'kind': 'source', 'value': value, 'unit': '项',
                'period': '2026上半年', 'population': '学院项目',
                'measure': '已开展项目数', 'value_type': 'actual',
                'source_ref': ref}
        item.update(other)
        return item

    def test_csv_blank_zero_percentage_and_fingerprint(self):
        source, ledger = self.ledger('a.csv', '项目,值,占比\n甲,0,0%\n乙,,20%\n')
        lookup = {x['id']: x for x in ledger['cells']}
        self.assertEqual(lookup['r000002c0002']['state'], 'zero')
        self.assertEqual(lookup['r000003c0002']['state'], 'blank')
        self.assertEqual(lookup['r000003c0003']['numeric_value'], 20.0)
        data, refs, cells, inputs = load_data_sources(['data.json'], self.base)
        self.assertEqual(refs['table:r000002c0002'], '0')
        self.assertEqual(len(inputs), 2)
        source.write_text('项目,值,占比\n甲,1,0%\n乙,,20%\n', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'differs from source'):
            load_data_sources(['data.json'], self.base)

    def test_docx_table_cells_keep_row_and_column_positions(self):
        docx = self.base / 'table.docx'
        ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        xml = f'''<w:document xmlns:w="{ns}"><w:body><w:tbl>
          <w:tr><w:tc><w:p><w:r><w:t>项目</w:t></w:r></w:p></w:tc>
                <w:tc><w:p><w:r><w:t>数量</w:t></w:r></w:p></w:tc></w:tr>
          <w:tr><w:tc><w:p><w:r><w:t>课程</w:t></w:r></w:p></w:tc>
                <w:tc><w:p><w:r><w:t>8</w:t></w:r></w:p></w:tc></w:tr>
          <w:tr><w:tc><w:p/></w:tc><w:tc><w:p/></w:tc></w:tr>
          </w:tbl></w:body></w:document>'''
        with ZipFile(docx, 'w') as z:
            z.writestr('word/document.xml', xml)
        ledger = extract_data(docx, 'wordtable')
        self.assertEqual(ledger['detail']['table_count'], 1)
        self.assertEqual([c['id'] for c in ledger['cells']],
                         ['t001r0001c001', 't001r0001c002',
                          't001r0002c001', 't001r0002c002',
                          't001r0003c001', 't001r0003c002'])
        self.assertEqual(ledger['cells'][3]['numeric_value'], 8)
        self.assertEqual(ledger['cells'][-1]['state'], 'blank')

    def test_xlsx_raw_formula_cache_and_visible_percent(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = '项目'
        ws['A1'] = 0.2
        ws['A1'].number_format = '0%'
        ws['A2'] = '=SUM(2,4)'
        ws['A3'] = 0
        source = self.base / 'table.xlsx'
        wb.save(source)
        ledger = extract_data(source, 'book')
        cells = {x['id']: x for x in ledger['cells']}
        self.assertEqual(cells['项目!A1']['raw'], 0.2)
        self.assertEqual(cells['项目!A1']['numeric_value'], 20)
        self.assertEqual(cells['项目!A2']['formula'], '=SUM(2,4)')
        self.assertEqual(cells['项目!A2']['state'], 'formula_uncached')
        self.assertEqual(cells['项目!A3']['state'], 'zero')
        cached = self.base / 'cached.xlsx'
        with ZipFile(source) as old, ZipFile(cached, 'w') as new:
            for name in old.namelist():
                content = old.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    self.assertIn(b'<f>SUM(2,4)</f><v></v>', content)
                    content = content.replace(b'<f>SUM(2,4)</f><v></v>',
                                              b'<f>SUM(2,4)</f><v>6</v>')
                new.writestr(name, content)
        cached_cell = {x['id']: x for x in extract_data(cached, 'book')['cells']}['项目!A2']
        self.assertEqual(cached_cell['cached_value'], 6)
        self.assertEqual(cached_cell['numeric_value'], 6)

    def test_metric_calculations_and_comparability(self):
        _, ledger = self.ledger('counts.csv', '2,6,8,10,0\n')
        cells = {'table:' + c['id']: c for c in ledger['cells']}
        a = self.metric('a', 2, 'table:r000001c0001', slice='A')
        b = self.metric('b', 6, 'table:r000001c0002', slice='B')
        total = self.metric('total', 8, 'table:r000001c0003', slice='all')
        target = self.metric('target', 10, 'table:r000001c0004', value_type='target', slice='all')
        metrics = [a, b, total, target,
                   {'id': 'sum', 'kind': 'derived', 'value': 8, 'unit': '项',
                    'period': '2026上半年', 'population': '学院项目', 'measure': '已开展项目数',
                    'value_type': 'actual', 'operation': 'sum', 'operands': ['a', 'b'],
                    'aggregation_basis': 'A/B 两类互斥'},
                   {'id': 'share', 'kind': 'derived', 'value': 25, 'unit': '%',
                    'period': '2026上半年', 'population': '学院项目', 'measure': 'A类项目占比',
                    'value_type': 'actual', 'operation': 'ratio_percent',
                    'operands': ['a', 'total'], 'denominator': 'total', 'relation': 'subset',
                    'decimals': 0},
                   {'id': 'progress', 'kind': 'derived', 'value': 80, 'unit': '%',
                    'period': '2026上半年', 'population': '学院项目', 'measure': '目标完成率',
                    'value_type': 'comparison', 'operation': 'target_progress_percent',
                    'operands': ['total', 'target'], 'denominator': 'target', 'decimals': 0}]
        result = evaluate_metrics(metrics, cells)
        self.assertEqual(result['share']['value'], 25)
        self.assertEqual(result['progress']['value'], 80)
        bad = [dict(x) for x in metrics]
        bad[1]['period'] = '2025上半年'
        with self.assertRaisesRegex(ValueError, 'incompatible period'):
            evaluate_metrics(bad, cells)
        bad = [dict(x) for x in metrics]
        bad[-1]['denominator'] = 'total'
        with self.assertRaisesRegex(ValueError, 'explicit target denominator'):
            evaluate_metrics(bad, cells)

    def test_missing_metric_is_distinct_from_zero_and_cannot_be_derived(self):
        _, ledger = self.ledger('missing.csv', ',0\n')
        cells = {'table:' + c['id']: c for c in ledger['cells']}
        missing = self.metric('missing', None, 'table:r000001c0001',
                              missing_reason='原表未填报')
        zero = self.metric('zero', 0, 'table:r000001c0002')
        self.assertIsNone(evaluate_metrics([missing, zero], cells)['missing']['value'])
        derived = {'id': 'bad', 'kind': 'derived', 'value': 0, 'unit': '项',
                   'period': '2026上半年', 'population': '学院项目',
                   'measure': '已开展项目数', 'value_type': 'actual',
                   'operation': 'sum', 'operands': ['missing', 'zero'],
                   'aggregation_basis': '两个互斥类别'}
        with self.assertRaisesRegex(ValueError, 'Missing operand'):
            evaluate_metrics([missing, zero, derived], cells)

    def test_period_growth_percentage_points_and_explicit_denominator(self):
        _, ledger = self.ledger('periods.csv', '10,15,40%,55%,0\n')
        cells = {'table:' + c['id']: c for c in ledger['cells']}
        previous = self.metric('previous', 10, 'table:r000001c0001', period='2025')
        current = self.metric('current', 15, 'table:r000001c0002', period='2026')
        p_previous = self.metric('p_previous', 40, 'table:r000001c0003',
                                 period='2025', unit='%', measure='参与率',
                                 denominator_description='2025年参与对象总数')
        p_current = self.metric('p_current', 55, 'table:r000001c0004',
                                period='2026', unit='%', measure='参与率',
                                denominator_description='2026年参与对象总数')
        growth = {'id': 'growth', 'kind': 'derived', 'value': 50, 'unit': '%',
                  'period': '2026', 'population': '学院项目', 'measure': '已开展项目数增长率',
                  'value_type': 'comparison', 'operation': 'growth_percent',
                  'operands': ['current', 'previous'], 'comparison_basis': '同口径年度项目数',
                  'decimals': 0}
        points = {'id': 'points', 'kind': 'derived', 'value': 15, 'unit': '百分点',
                  'period': '2026', 'population': '学院项目', 'measure': '参与率变化',
                  'value_type': 'comparison', 'operation': 'percentage_point_change',
                  'operands': ['p_current', 'p_previous'],
                  'comparison_basis': '同一参与对象定义', 'decimals': 0}
        data = [previous, current, p_previous, p_current, growth, points]
        result = evaluate_metrics(data, cells)
        self.assertEqual(result['growth']['value'], 50)
        self.assertEqual(result['points']['value'], 15)
        wrong = [dict(x) for x in data]
        wrong[1]['value_type'] = 'estimate'
        with self.assertRaisesRegex(ValueError, 'incompatible value_type'):
            evaluate_metrics(wrong, cells)
        wrong = [dict(x) for x in data]
        wrong[3].pop('denominator_description')
        with self.assertRaisesRegex(ValueError, 'requires denominator'):
            evaluate_metrics(wrong, cells)

    def test_table_count_uses_contiguous_range_and_status_prefix(self):
        _, ledger = self.ledger('status.csv', '课程,状态\n甲,已投入教学使用\n乙,待开展\n丙,已投入教学使用；附评审说明\n')
        cells = {'table:' + c['id']: c for c in ledger['cells']}
        status = {'id': 'active', 'kind': 'table_count', 'value': 2, 'unit': '门',
                  'period': '2026细则', 'population': '列示课程', 'measure': '投入教学使用课程数',
                  'value_type': 'actual', 'count_basis': '状态以已投入教学使用开头',
                  'query': {'source_id': 'table', 'format': 'csv', 'column': 2,
                            'first_row': 2, 'last_row': 4,
                            'starts_with': ['已投入教学使用']}}
        self.assertEqual(evaluate_metrics([status], cells)['active']['value'], 2)
        bad = dict(status)
        bad['query'] = dict(status['query'], last_row=3)
        with self.assertRaisesRegex(ValueError, 'Incorrect table count'):
            evaluate_metrics([bad], cells)
        bad['query'] = dict(status['query'], last_row=5)
        with self.assertRaisesRegex(ValueError, 'missing cells'):
            evaluate_metrics([bad], cells)

    def test_metric_id_can_bind_legacy_page_and_check_data_cli(self):
        _, ledger = self.ledger('counts.csv', '项目,值\n课程,8\n')
        ref = extract_source(TEMPLATE, 'reference', 'evidence')
        write_json(self.base / 'reference.json', ref)
        inv = inventory()
        page = {'id': 'opening', 'template_slide': 1, 'purpose': '说明课程项目数',
                'slots': {}, 'pictures': {}, 'notes_refs': ['reference:s001']}
        for slot in inv['slides'][0]['slots']:
            if slot['text'] and slot['name'] != 'page-num' and slot['type'] != 'grpSp':
                page['slots'][slot['key']] = {'action': 'keep', 'kind': 'source',
                                               'sources': ['reference:s001']}
        page['slots']['6'] = {'action': 'replace', 'kind': 'source', 'text': '课程8项',
                               'sources': ['table:r000002c0002'], 'facts': ['courses']}
        plan = {'schema_version': 1, 'template_sha256': file_sha(TEMPLATE),
                'sources': ['reference.json'], 'data_sources': ['data.json'],
                'metrics': [self.metric('courses', 8, 'table:r000002c0002')],
                'facts': [], 'pages': [page], 'gaps': []}
        write_json(self.base / 'plan.json', plan)
        sources, refs, inputs, gaps = validate_plan(plan, self.base)
        self.assertEqual(gaps, [])
        self.assertEqual(evaluate_plan_facts(plan, refs, sources)['courses']['value'], 8)
        cli = Path(__file__).resolve().parents[1] / 'scripts' / 'personal_ppt.py'
        run = subprocess.run([sys.executable, '-X', 'utf8', str(cli), 'check-data',
                              str(self.base / 'plan.json')], capture_output=True,
                             text=True, encoding='utf-8')
        self.assertEqual(run.returncode, 0, run.stderr)
        output = json.loads(run.stdout)
        self.assertEqual(output['data_sources'], 1)
        self.assertEqual(output['metrics'][0]['id'], 'courses')


if __name__ == '__main__':
    unittest.main()
