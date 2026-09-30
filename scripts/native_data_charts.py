"""Editable, workbook-backed chart parts for the personal native-PPTX builder.

The reference slide remains the canvas.  A small temporary presentation supplies
valid Office chart parts, which are transplanted as native DrawingML objects.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import math
import posixpath as pp
import re

from lxml import etree as E

from native_pptx import (NS, EMU, q, xml, serialize, shapes, bounds, relationships,
                         add_rel, relpath, target, content_type)

CHART_TYPES = {'trend', 'comparison', 'ranking', 'composition', 'target_progress'}


def validate_data_chart(spec, page, ss, facts):
    """Reject ambiguous or unsupported visual claims before touching a package."""
    from native_pptx import text_of
    if not isinstance(spec, dict) or spec.get('type') not in CHART_TYPES:
        raise ValueError('data_chart needs a supported type')
    categories = spec.get('categories')
    series = spec.get('series')
    if not isinstance(categories, list) or not 1 <= len(categories) <= 12 or any(
        not isinstance(c, str) or not c.strip() or len(c) > 24 for c in categories
    ) or len(set(categories)) != len(categories):
        raise ValueError('data_chart needs 1–12 distinct, concise categories')
    if not isinstance(series, list) or not 1 <= len(series) <= 3:
        raise ValueError('data_chart needs 1–3 series')
    if spec['type'] in ('ranking', 'composition') and len(series) != 1:
        raise ValueError('ranking/composition need one series')
    if spec['type']=='composition' and len(categories)>6:
        raise ValueError('composition needs at most six legible slices; group or split categories')
    if spec['type'] == 'target_progress' and len(series) != 2:
        raise ValueError('target_progress needs actual and target series')
    if spec['type'] in ('trend', 'comparison') and len(categories) < 2:
        raise ValueError('trend/comparison need multiple categories')
    all_facts=[]
    for s in series:
        if not isinstance(s, dict) or not isinstance(s.get('name'), str) or not s['name'].strip():
            raise ValueError('data_chart series needs a name')
        ids=s.get('fact_ids')
        if not isinstance(ids, list) or len(ids) != len(categories):
            raise ValueError('data_chart series needs one metric per category')
        for fid in ids:
            if fid not in facts or facts[fid].get('value') is None:
                raise ValueError('Unknown or missing data_chart metric: '+str(fid))
            value=facts[fid]['value']
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
                raise ValueError('data_chart metrics must be finite numbers')
            all_facts.append(facts[fid])
    if len({f['unit'] for f in all_facts}) != 1:
        raise ValueError('data_chart mixes measurement units')
    for dimension in ('population','measure'):
        if any(dimension in f for f in all_facts) and (
            any(not f.get(dimension) for f in all_facts) or
            len({f[dimension] for f in all_facts}) != 1
        ):
            raise ValueError('data_chart mixes '+dimension+' definitions')
    if spec['type']=='trend':
        for s in series:
            periods=[facts[fid]['period'] for fid in s['fact_ids']]
            if len(set(periods))!=len(periods):
                raise ValueError('trend categories must map to distinct metric periods')
    if spec['type'] == 'composition':
        values=[facts[fid]['value'] for fid in series[0]['fact_ids']]
        if any(v < 0 for v in values) or sum(values) <= 0:
            raise ValueError('composition needs nonnegative values with positive sum')
        if len({f['period'] for f in all_facts}) != 1:
            raise ValueError('composition needs one time period')
    if spec['type'] in ('comparison','ranking','target_progress') and any(
        facts[fid]['value'] < 0 for s in series for fid in s['fact_ids']
    ):
        raise ValueError('comparison/ranking/progress need nonnegative values')
    if spec['type'] == 'ranking':
        values=[facts[fid]['value'] for fid in series[0]['fact_ids']]
        if values != sorted(values, reverse=True):
            raise ValueError('ranking must be sorted descending')
    if spec['type'] in ('ranking','target_progress') and len({f['period'] for f in all_facts})!=1:
        raise ValueError('ranking/progress need a common time period')
    if spec['type']=='target_progress' and any('value_type' in f for f in all_facts):
        if ({facts[fid].get('value_type') for fid in series[0]['fact_ids']}!={'actual'} or
            {facts[fid].get('value_type') for fid in series[1]['fact_ids']}!={'target'}):
            raise ValueError('target_progress needs actual then target metrics')
    box=spec.get('bounds_pt')
    if not isinstance(box,list) or len(box)!=4 or any(
        isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) for v in box
    ) or box[0]<0 or box[1]<0 or box[2]<80 or box[3]<80 or box[0]+box[2]>960 or box[1]+box[3]>540:
        raise ValueError('Invalid data_chart bounds')
    if not isinstance(spec.get('source_label'),str) or not spec['source_label'].strip():
        raise ValueError('data_chart needs a visible source label')
    if not isinstance(spec.get('interpretation'),str) or not spec['interpretation'].strip():
        raise ValueError('data_chart needs an interpretation in the content plan')
    removed=spec.get('replace_shapes',[])
    if not isinstance(removed,list) or len(set(removed))!=len(removed):
        raise ValueError('data_chart replace_shapes needs unique object IDs')
    chart_box=box
    for key in removed:
        if key not in ss:
            raise ValueError('Unknown data_chart replacement shape '+str(key))
        node=ss[key]
        if node.tag in (q('p:pic'),q('p:grpSp'),q('p:graphicFrame')) or key=='6' or node.find('.//p:cNvPr',NS).get('name') in ('logo','page-num'):
            raise ValueError('data_chart may not replace identity, title, photos or groups')
        if text_of(node) and page.get('slots',{}).get(key,{}).get('action')!='remove':
            raise ValueError('Historical chart-region text needs explicit removal binding: '+key)
        b=bounds(node)
        if b and not (b[0]<chart_box[0]+chart_box[2] and b[0]+b[2]>chart_box[0] and b[1]<chart_box[1]+chart_box[3] and b[1]+b[3]>chart_box[1]):
            raise ValueError('data_chart replacement shape lies outside chart bounds: '+key)
    covered=set(removed)|set(page.get('remove_shapes',[]))
    for key,node in ss.items():
        b=bounds(node)
        if b and key not in covered and b[0]<chart_box[0]+chart_box[2] and b[0]+b[2]>chart_box[0] and b[1]<chart_box[1]+chart_box[3] and b[1]+b[3]>chart_box[1]:
            raise ValueError('Historical shape overlaps new data_chart: '+key)
    if not spec.get('reason'):
        raise ValueError('data_chart shape replacement needs reason')


def _donor(spec, facts):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    data=CategoryChartData();data.categories=spec['categories']
    for s in spec['series']:
        data.add_series(s['name'],[facts[fid]['value'] for fid in s['fact_ids']])
    typ={'trend':XL_CHART_TYPE.LINE_MARKERS,
         'comparison':XL_CHART_TYPE.COLUMN_CLUSTERED,
         'ranking':XL_CHART_TYPE.BAR_CLUSTERED,
         'composition':XL_CHART_TYPE.DOUGHNUT,
         'target_progress':XL_CHART_TYPE.BAR_CLUSTERED}[spec['type']]
    x,y,w,h=spec['bounds_pt']
    frame=slide.shapes.add_chart(typ,Pt(x),Pt(y),Pt(w),Pt(h),data)
    chart=frame.chart;chart.has_title=False
    chart.has_legend=len(spec['series'])>1 or spec['type']=='composition'
    chart.chart_style=13
    chart.font.name='MiSans Normal';chart.font.size=Pt(11)
    chart.font.color.rgb=RGBColor(24,49,88)
    if spec['type']!='composition':
        for axis in (chart.category_axis,chart.value_axis):
            axis.tick_labels.font.name='MiSans Normal'
            axis.tick_labels.font.size=Pt(10)
            axis.tick_labels.font.color.rgb=RGBColor(54,72,98)
            axis.format.line.color.rgb=RGBColor(178,189,204)
            axis.format.line.width=Pt(.6)
        chart.value_axis.has_major_gridlines=True
        chart.value_axis.major_gridlines.format.line.color.rgb=RGBColor(217,224,233)
        chart.value_axis.major_gridlines.format.line.width=Pt(.5)
        chart.value_axis.has_title=True
        chart.value_axis.axis_title.text_frame.text=facts[spec['series'][0]['fact_ids'][0]]['unit']
        chart.value_axis.axis_title.text_frame.paragraphs[0].font.name='MiSans Normal'
        chart.value_axis.axis_title.text_frame.paragraphs[0].font.size=Pt(9)
    if chart.has_legend:
        chart.legend.position=(XL_LEGEND_POSITION.RIGHT if spec['type']=='composition'
                               else XL_LEGEND_POSITION.BOTTOM)
        chart.legend.include_in_layout=False
        chart.legend.font.name='MiSans Normal';chart.legend.font.size=Pt(10)
    palette=[RGBColor(24,49,88),RGBColor(237,129,35),RGBColor(112,136,171)]
    for i,series in enumerate(chart.series):
        series.format.fill.solid();series.format.fill.fore_color.rgb=palette[i]
        try:
            series.format.line.color.rgb=palette[i]
            series.format.line.width=Pt(2.25)
        except (AttributeError,ValueError):pass
    if spec['type']=='trend':
        chart.plots[0].has_data_labels=False
        for i,series in enumerate(chart.series):
            series.marker.style=XL_MARKER_STYLE.CIRCLE
            series.marker.size=7
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb=palette[i]
            series.marker.format.line.color.rgb=RGBColor(255,255,255)
    if spec['type'] in ('comparison','ranking','target_progress'):
        chart.plots[0].gap_width=70
    if spec['type'] in ('ranking','target_progress'):
        chart.category_axis.reverse_order=True
    if spec['type']=='composition':
        chart.plots[0].hole_size=64
        colors=[RGBColor(24,49,88),RGBColor(237,129,35),RGBColor(78,119,169),
                RGBColor(146,167,194),RGBColor(193,204,217),RGBColor(233,186,146)]
        for i,point in enumerate(chart.series[0].points):
            point.format.fill.solid();point.format.fill.fore_color.rgb=colors[i%len(colors)]
    out=BytesIO();prs.save(out)
    with ZipFile(BytesIO(out.getvalue())) as z:
        return {n:z.read(n) for n in z.namelist() if not n.endswith('/')}


def _normalize_data_sheet(chart_bytes, workbook_bytes):
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(workbook_bytes))
    if len(wb.worksheets)!=1:raise ValueError('Donor chart workbook topology changed')
    old=wb.active.title;wb.active.title='Data'
    out=BytesIO();wb.save(out)
    chart=xml(chart_bytes)
    for f in chart.findall('.//c:f',NS):
        if f.text:f.text=f.text.replace(old+'!',"'Data'!").replace("'"+old+"'!","'Data'!")
    return serialize(chart),out.getvalue()


def add_data_chart(parts, slide_root, slide_rels, slide_part, index, spec, facts):
    """Attach chart and workbook and return an auditable chart manifest record."""
    donor=_donor(spec,facts)
    donor_slide=xml(donor['ppt/slides/slide1.xml'])
    frame=deepcopy(donor_slide.find('.//p:graphicFrame',NS))
    if frame is None:raise ValueError('Donor chart frame missing')
    existing={int(k.split('/')[-1]) for k in shapes(slide_root) if k.split('/')[-1].isdigit()}
    new_id=max(existing or {0})+1
    frame.find('.//p:cNvPr',NS).set('id',str(new_id))
    frame.find('.//p:cNvPr',NS).set('name','Seqara evidence chart '+str(index))
    chart_ref=frame.find('.//c:chart',NS)
    if chart_ref is None:raise ValueError('Donor graphic frame is not a chart')
    donor_slide_rels=relationships(donor,'ppt/slides/slide1.xml')
    donor_chart_rel=next((r for r in donor_slide_rels if r.get('Id')==chart_ref.get(q('r:id'))),None)
    if donor_chart_rel is None:raise ValueError('Donor chart relationship missing')
    donor_chart_part=target('ppt/slides/slide1.xml',donor_chart_rel.get('Target'))
    donor_chart_rels=relationships(donor,donor_chart_part)
    wb_rel=next((r for r in donor_chart_rels if r.get('Type').endswith('/package')),None)
    if wb_rel is None:raise ValueError('Donor chart workbook missing')
    donor_workbook=target(donor_chart_part,wb_rel.get('Target'))
    chart_part=f'ppt/charts/seqara-data-{index}.xml'
    workbook=f'ppt/embeddings/seqara-data-{index}.xlsx'
    chart_bytes,wb_bytes=_normalize_data_sheet(donor[donor_chart_part],donor[donor_workbook])
    parts[chart_part]=chart_bytes;parts[workbook]=wb_bytes
    chart_rels=relationships({},chart_part)
    add_rel(chart_rels,'rIdDataWorkbook','package',pp.relpath(workbook,pp.dirname(chart_part)))
    chart=xml(chart_bytes);external=chart.find('c:externalData',NS)
    if external is None:raise ValueError('Donor chart data binding missing')
    external.set(q('r:id'),'rIdDataWorkbook')
    parts[chart_part]=serialize(chart)
    parts[relpath(chart_part)]=serialize(chart_rels)
    content_type(parts,chart_part,'application/vnd.openxmlformats-officedocument.drawingml.chart+xml')
    content_type(parts,workbook,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    used={r.get('Id') for r in slide_rels}
    rid='rIdSeqaraChart'+str(index)
    if rid in used:raise ValueError('Chart relationship ID collision')
    add_rel(slide_rels,rid,'chart',pp.relpath(chart_part,pp.dirname(slide_part)))
    chart_ref.set(q('r:id'),rid)
    slide_root.find('p:cSld/p:spTree',NS).append(frame)
    return {'type':spec['type'],'chart_part':chart_part,'workbook_part':workbook,
            'categories':spec['categories'],'series':[{'name':s['name'],'fact_ids':s['fact_ids']} for s in spec['series']],
            'unit':facts[spec['series'][0]['fact_ids'][0]]['unit'],
            'source_label':spec['source_label'],'interpretation':spec['interpretation']}
