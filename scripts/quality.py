"""Checks the exported artifact. No synthetic visual approval."""
from pathlib import Path
from io import BytesIO
import importlib.util
import math
import platform
import re
from native_pptx import *

# These require a whole-deck judgment after the individual slide inspections.
# They deliberately cannot be inferred from shape counts or render success.
DECK_REVIEW_KEYS=(
    'audience_fit',
    'title_argument_order',
    'visual_semantics',
    'density_transitions',
    'notes_support',
)

def doctor():
    deps={m:importlib.util.find_spec(m) is not None for m in ('lxml','PIL','openpyxl','pptx','win32com')}
    fonts=[];ppt=False
    if platform.system()=='Windows':
        import winreg
        for hive in (winreg.HKEY_CURRENT_USER,winreg.HKEY_LOCAL_MACHINE):
            try:
                with winreg.OpenKey(hive,r'SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts') as k:
                    for i in range(winreg.QueryInfoKey(k)[1]):
                        name,value,_=winreg.EnumValue(k,i)
                        if 'misans' in name.lower():fonts.append({'name':name,'file':value})
            except OSError:pass
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,r'PowerPoint.Application\CLSID'):ppt=True
        except OSError:pass
    ok=all(deps.values()) and ppt and any('misans normal' in f['name'].lower() for f in fonts)
    return {'ready_for_verified_delivery':ok,'platform':platform.system(),'dependencies':deps,'powerpoint_registered':ppt,
            'fonts':fonts,'template_sha256':file_sha(TEMPLATE),
            'font_policy':'MiSans Normal required; no silent substitution','wysiwyg_engine':'Microsoft PowerPoint'}

def validate_chart(parts,chart_part):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    c=xml(parts[chart_part]);rels={r.get('Id'):r for r in relationships(parts,chart_part)}
    ext=c.find('c:externalData',NS)
    if ext is None:raise ValueError('Native chart has no workbook binding')
    rel=rels.get(ext.get(q('r:id')))
    if rel is None or rel.get('TargetMode')=='External' or not rel.get('Type').endswith('/package'):raise ValueError('Chart data link is not embedded')
    wb=load_workbook(BytesIO(parts[target(chart_part,rel.get('Target'))]),data_only=False)
    if 'Data' not in wb.sheetnames:raise ValueError('Chart workbook lacks Data sheet')
    ws=wb['Data']
    def cells(formula):
        if not formula or not re.fullmatch(r"(?:'Data'|Data)!\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?",formula):
            raise ValueError('Chart formula is not an embedded Data-sheet range')
        coord=formula.split('!',1)[1].replace('$','')
        min_col,min_row,max_col,max_row=range_boundaries(coord)
        return [ws.cell(row,col).value for row in range(min_row,max_row+1)
                for col in range(min_col,max_col+1)]
    def read_ref(ser,kind,numeric=False):
        node=ser.find('c:'+kind,NS)
        if node is None:raise ValueError('Chart series lacks '+kind)
        ref=node.find('c:numRef' if numeric else 'c:strRef',NS)
        if ref is None:raise ValueError('Chart series uses unsupported literal data')
        f=ref.find('c:f',NS)
        if f is None:raise ValueError('Chart series has no workbook formula')
        actual=cells(f.text)
        cache=ref.find('c:numCache' if numeric else 'c:strCache',NS)
        if cache is None:raise ValueError('Chart series lacks cache')
        cached=[(int(p.get('idx')),p.find('c:v',NS).text if p.find('c:v',NS) is not None else '')
                for p in cache.findall('c:pt',NS)]
        cached=[v for _,v in sorted(cached)]
        if len(actual)!=len(cached):raise ValueError('Chart cache length differs from workbook')
        if numeric:
            if any(a is None or not math.isclose(float(a),float(b),rel_tol=0,abs_tol=1e-9)
                   for a,b in zip(actual,cached)):
                raise ValueError('Chart numeric cache and workbook disagree')
            return [float(v) for v in cached]
        if [str(v) for v in actual]!=cached:
            raise ValueError('Chart text cache and workbook disagree')
        return cached
    records=[]
    for ser in c.findall('.//c:ser',NS):
        name=read_ref(ser,'tx')
        cats=read_ref(ser,'cat')
        vals=read_ref(ser,'val',True)
        if len(name)!=1 or len(cats)!=len(vals):raise ValueError('Chart series ranges disagree')
        records.append({'name':name[0],'categories':cats,'values':vals})
    if not records:raise ValueError('Chart has no series')
    if any(r['categories']!=records[0]['categories'] for r in records):
        raise ValueError('Chart series category ranges disagree')
    return {'part':chart_part,'categories':records[0]['categories'],
            'values':records[0]['values'],'series':records,'embedded_workbook':True}

def structure_check(directory):
    d=Path(directory);manifest=read_json(d/'manifest.json');pptx=d/'presentation.pptx';parts=package(pptx);errors=[];charts=[]
    if file_sha(pptx)!=manifest['pptx_sha256']:errors.append('Output fingerprint changed')
    for source in manifest['inputs']:
        p=Path(source['path'])
        if not p.exists() or file_sha(p)!=source['sha256']:errors.append('Input changed or missing: '+str(p))
    slideparts=ordered_slides(parts)
    if len(slideparts)!=len(manifest['pages']):errors.append('Slide count mismatch')
    for expected,part in zip(manifest['pages'],slideparts):
        if sha(parts[part])!=expected['slide_xml_sha256']:errors.append(f'Slide XML changed: {expected["index"]}')
        root=xml(parts[part]);ss=shapes(root)
        for key,obj in expected['objects'].items():
            if key not in ss or style_signature(ss[key])!=obj['signature'] or text_of(ss[key])!=obj['text']:
                errors.append(f'Slot format/content differs: {expected["index"]}/{key}')
        rr=relationships(parts,part);notes=[r for r in rr if r.get('Type').endswith('/notesSlide')]
        if len(notes)!=1:errors.append('Each page requires notes')
        else:
            note=xml(parts[target(part,notes[0].get('Target'))]);txt=text_of(note)
            if sha(txt.encode())!=expected['notes_sha256']:errors.append('Notes content mismatch on '+str(expected['index']))
    pages_by_id={p['id']:p for p in manifest['pages']}
    for binding in manifest.get('media_bindings',[]):
        page=pages_by_id.get(binding['page_id'])
        if page is None:
            errors.append('Media binding has missing page '+binding['page_id']);continue
        part=page['part'];ss=shapes(xml(parts[part]));node=ss.get(binding['shape_key'])
        if node is None or node.tag!=q('p:pic'):
            errors.append('Media binding has missing picture '+binding['asset_id']);continue
        blip=node.find('.//a:blip',NS)
        rels={r.get('Id'):r for r in relationships(parts,part)}
        rel=rels.get(blip.get(q('r:embed'))) if blip is not None else None
        media=target(part,rel.get('Target')) if rel is not None else None
        if media not in parts or sha(parts[media])!=binding['media_sha256']:
            errors.append('Media output differs from DOCX asset '+binding['asset_id'])
        crop=node.find('.//a:srcRect',NS)
        actual=dict(crop.attrib) if crop is not None else {}
        if binding.get('crop') is not None and {k:str(v) for k,v in binding['crop'].items()}!=actual:
            errors.append('Media crop differs from plan '+binding['asset_id'])
    for name in list(parts):
        if name.endswith('.rels'):
            owner='' if name=='_rels/.rels' else pp.join(pp.dirname(pp.dirname(name)),pp.basename(name)[:-5])
            for r in xml(parts[name]):
                if r.get('TargetMode')=='External':
                    if not r.get('Target'):errors.append('Empty external relationship')
                elif target(owner,r.get('Target')) not in parts:errors.append('Broken relationship '+name)
        if re.fullmatch(r'ppt/charts/[^/]+\.xml',name):
            try:charts.append(validate_chart(parts,name))
            except Exception as exc:errors.append(str(exc))
    chart_by_part={c['part']:c for c in charts}
    for binding in manifest.get('data_charts',[]):
        chart=chart_by_part.get(binding['chart_part'])
        if chart is None:
            errors.append('Data chart missing: '+binding['chart_part']);continue
        if chart['categories']!=binding['categories'] or [s['name'] for s in chart['series']]!=[s['name'] for s in binding['series']]:
            errors.append('Data chart categories/series differ from plan: '+binding['chart_part'])
    return {'errors':errors,'charts':charts,'pages':len(slideparts),'pptx_sha256':file_sha(pptx)}

def contact_sheets(render_dir, out, per_sheet=6):
    from PIL import Image,ImageDraw
    files=sorted(Path(render_dir).glob('slide-*.png'));results=[]
    for offset in range(0,len(files),per_sheet):
        count=min(per_sheet,len(files)-offset)
        canvas=Image.new('RGB',(1280,math.ceil(count/2)*384),(231,233,237));draw=ImageDraw.Draw(canvas)
        for j,p in enumerate(files[offset:offset+per_sheet]):
            with Image.open(p) as im:
                im=im.convert('RGB');im.thumbnail((624,351));x=(j%2)*640+8;y=(j//2)*384+24
                canvas.paste(im,(x,y));draw.text((x,y-20),p.stem,fill=(20,35,60))
        path=Path(out)/f'contact-{offset//per_sheet+1:02d}.jpg';canvas.save(path,quality=91);results.append(str(path))
    return results

def render_powerpoint(pptx, directory):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize();app=None;deck=None;existing=False
    d=Path(directory);d.mkdir(parents=True,exist_ok=True);render=d/'render';render.mkdir(exist_ok=True)
    results=[];pdf=d/'presentation.pdf'
    try:
        try:
            win32com.client.GetActiveObject('PowerPoint.Application');existing=True
        except Exception:pass
        app=win32com.client.DispatchEx('PowerPoint.Application')
        deck=app.Presentations.Open(str(Path(pptx).resolve()),True,False,False)
        def inspect(collection, prefix=''):
            measured=[]
            for i in range(1,collection.Count+1):
                s=collection.Item(i);key=prefix+str(s.Id)
                if s.Type==6:
                    measured.extend(inspect(s.GroupItems,key+'/'));continue
                if s.HasTextFrame and s.TextFrame.HasText:
                    frame=s.TextFrame2;tr=frame.TextRange
                    width=float(s.Width-frame.MarginLeft-frame.MarginRight);height=float(s.Height-frame.MarginTop-frame.MarginBottom)
                    usedw=float(tr.BoundWidth);usedh=float(tr.BoundHeight)
                    text=str(tr.Text)
                    # Empty bullets/whitespace are excluded from measurement, not actual content.
                    if text.strip():
                        fonts=set();sizes=set()
                        try:
                            for k in range(1,tr.Runs().Count+1):
                                run=tr.Runs(k,1);fonts.add(str(run.Font.Name));sizes.add(float(run.Font.Size))
                        except Exception:fonts.add(str(tr.Font.Name))
                        measured.append({'key':key,'name':s.Name,'text':text,'width':width,'height':height,
                                         'used_width':round(usedw,3),'used_height':round(usedh,3),
                                         'overflow_x':round(max(0,usedw-width),3),'overflow_y':round(max(0,usedh-height),3),
                                         'fonts':sorted(fonts),'font_sizes':sorted(sizes)})
            return measured
        for i in range(1,deck.Slides.Count+1):
            slide=deck.Slides.Item(i);png=render/f'slide-{i:02d}.png';slide.Export(str(png),'PNG',1600,900)
            measurements=inspect(slide.Shapes)
            results.append({'index':i,'render':str(png),'render_sha256':file_sha(png),'measurements':measurements})
        deck.SaveAs(str(pdf),32)
        return {'engine':'Microsoft PowerPoint','engine_version':str(app.Version),'slides':results,'pdf':str(pdf),'pdf_sha256':file_sha(pdf),
                'contacts':contact_sheets(render,d),'render_size':[1600,900]}
    finally:
        if deck is not None:deck.Close()
        # PowerPoint is a singleton on some machines: never terminate an existing user's app.
        if app is not None and not existing:
            try:
                if app.Presentations.Count==0:app.Quit()
            except Exception:pass
        pythoncom.CoUninitialize()

def verify(directory, render=True):
    d=Path(directory).resolve();report=structure_check(d);env=doctor();report['environment']=env
    report['status']='structure_checked';report['warnings']=[]
    if not env['ready_for_verified_delivery']:report['errors'].append('Environment incomplete: MiSans Normal, PowerPoint and dependencies required')
    if render and env['ready_for_verified_delivery'] and not report['errors']:
        try:
            report['rendering']=render_powerpoint(d/'presentation.pptx',d)
            baseline_path=ROOT/'assets/reference-metrics.json'
            baseline=read_json(baseline_path) if baseline_path.exists() else {}
            original=package(TEMPLATE);original_order=ordered_slides(original);manifest=read_json(d/'manifest.json')
            calibrated=(baseline.get('template_sha256')==file_sha(TEMPLATE) and baseline.get('engine_version')==report['rendering']['engine_version'])
            for slide in report['rendering']['slides']:
                expected=manifest['pages'][slide['index']-1];ref_number=expected['template_slide']
                ref_shapes=shapes(xml(original[original_order[ref_number-1]]))
                measurements={x['key']:x for x in baseline.get('slides',{}).get(str(ref_number),[])} if calibrated else {}
                for item in slide['measurements']:
                    if item['overflow_x']>1 or item['overflow_y']>1:
                        key=item['key'];m=measurements.get(key);old=ref_shapes.get(key);current=expected['objects'].get(key)
                        same=(m is not None and old is not None and current is not None and
                              text_of(old)==current['text'] and style_signature(old)==current['signature'] and
                              abs(item['overflow_x']-m['overflow_x'])<.1 and abs(item['overflow_y']-m['overflow_y'])<.1)
                        if same:
                            item['reference_inherited_overshoot']=True
                            report['warnings'].append(f'Reference text-frame overshoot page {slide["index"]} shape {key}; exact unchanged object, visually inspect clipping/adjacent content')
                        else:
                            report['errors'].append(f'Overflow page {slide["index"]} shape {item["key"]}: {item["overflow_x"]}/{item["overflow_y"]} pt')
                    if any(f and f!='MiSans Normal' for f in item['fonts']):
                        report['warnings'].append(f'Inspect font page {slide["index"]}: {item["fonts"]}')
            report['status']='rendered_pending_visual_review' if not report['errors'] else 'failed'
        except Exception as exc:report['errors'].append('PowerPoint rendering failed: '+str(exc));report['status']='failed'
    report['manifest_sha256']=file_sha(d/'manifest.json');write_json(d/'qa.json',report)
    return {'status':report['status'],'errors':report['errors'],'warnings':report['warnings'],'report':str(d/'qa.json')}

def review_init(directory):
    d=Path(directory);qa=read_json(d/'qa.json');manifest=read_json(d/'manifest.json')
    if 'rendering' not in qa:raise ValueError('Render before reviewing')
    result={'pptx_sha256':qa['pptx_sha256'],'qa_sha256':file_sha(d/'qa.json'),'reviewer':'',
            'method':'Inspect every full-resolution exported page and sources; compare representative pages with the reference, then review the full title sequence, diagram meanings, density, transitions and speaker notes as one talk.',
            'pages':[{'index':s['index'],'render_sha256':s['render_sha256'],'status':'pending',
                      'checks':{'layout':False,'typography':False,'evidence':False,'content':False,'editable':False},
                      'observations':'','issues':[]} for s in qa['rendering']['slides']],
            'deck_review':{'checks':{key:False for key in DECK_REVIEW_KEYS},
                           'observations':{key:'' for key in DECK_REVIEW_KEYS},
                           'overall_observations':''},
            'resolved_warnings':[],'material_limitations':[]}
    p=d/'visual-review.json'
    if p.exists():raise ValueError('Review exists; do not overwrite a human/agent review')
    write_json(p,result);return {'review':str(p),'pages':len(manifest['pages'])}

def finalize(directory):
    d=Path(directory);manifest=read_json(d/'manifest.json');qa=read_json(d/'qa.json');review=read_json(d/'visual-review.json')
    errors=structure_check(d)['errors']
    if qa.get('errors') or qa.get('status')!='rendered_pending_visual_review':errors.append('Export QA not passed')
    if qa.get('manifest_sha256')!=file_sha(d/'manifest.json'):errors.append('QA stale after manifest change')
    if review.get('pptx_sha256')!=file_sha(d/'presentation.pptx') or review.get('qa_sha256')!=file_sha(d/'qa.json'):errors.append('Visual review is stale')
    if not review.get('reviewer'):errors.append('Reviewer not recorded')
    pages=review.get('pages',[])
    if [p['index'] for p in pages]!=list(range(1,len(manifest['pages'])+1)):errors.append('Visual review must cover every page once')
    if [p['index'] for p in qa.get('rendering',{}).get('slides',[])]!=list(range(1,len(manifest['pages'])+1)):errors.append('Rendering must cover every page once')
    for p,r in zip(pages,qa.get('rendering',{}).get('slides',[])):
        if p.get('status')!='approved' or p.get('issues') or not p.get('observations') or set(p.get('checks',{}))!={'layout','typography','evidence','content','editable'} or not all(p['checks'].values()):errors.append('Incomplete review page '+str(p['index']))
        if not Path(r['render']).exists() or p.get('render_sha256')!=file_sha(r['render']):errors.append('Render fingerprint mismatch')
    # Old review files remain valid.  A newly initialized review must also
    # assess the deck as one argument, including the meaning of native icons.
    if 'deck_review' in review:
        deck_review=review['deck_review']
        checks=deck_review.get('checks') if isinstance(deck_review,dict) else None
        observations=deck_review.get('observations') if isinstance(deck_review,dict) else None
        overall=deck_review.get('overall_observations') if isinstance(deck_review,dict) else None
        if (not isinstance(checks,dict) or set(checks)!=set(DECK_REVIEW_KEYS) or
            any(checks[key] is not True for key in DECK_REVIEW_KEYS)):
            errors.append('Incomplete whole-deck review checks')
        if (not isinstance(observations,dict) or set(observations)!=set(DECK_REVIEW_KEYS) or
            any(not isinstance(observations[key],str) or not observations[key].strip()
                for key in DECK_REVIEW_KEYS) or
            not isinstance(overall,str) or not overall.strip()):
            errors.append('Incomplete whole-deck review observations')
    if set(qa.get('warnings',[]))-set(review.get('resolved_warnings',[])):errors.append('Unresolved QA warnings')
    if any(g.get('blocking',True) for g in manifest['gaps']):errors.append('Blocking material gaps remain')
    pdf=qa.get('rendering',{}).get('pdf')
    if not pdf or not Path(pdf).exists() or file_sha(pdf)!=qa['rendering']['pdf_sha256']:errors.append('PDF missing or changed')
    if errors:raise ValueError('; '.join(errors))
    delivery={'status':'verified','revision':manifest['revision'],'pptx_sha256':file_sha(d/'presentation.pptx'),
              'qa_sha256':file_sha(d/'qa.json'),'review_sha256':file_sha(d/'visual-review.json'),
              'pptx':str(d/'presentation.pptx'),'pdf':pdf,'limitations':review.get('material_limitations',[])}
    write_json(d/'delivery.json',delivery);return delivery

def compare_images(reference_dir, output_dir, mapping):
    from PIL import Image,ImageChops,ImageStat
    rows=[]
    for source,dest in mapping:
        with Image.open(Path(reference_dir)/f'slide-{source:02d}.png') as a,Image.open(Path(output_dir)/f'slide-{dest:02d}.png') as b:
            if a.size!=b.size:raise ValueError('Render both decks at same resolution')
            diff=ImageChops.difference(a.convert('RGB'),b.convert('RGB'));stats=ImageStat.Stat(diff)
            rows.append({'reference_page':source,'output_page':dest,'mean_absolute_channel_error':round(sum(stats.mean)/3,6),
                         'changed_bbox':diff.getbbox(),'note':'Page-number corrections are expected; inspect differences visually.'})
    return rows
